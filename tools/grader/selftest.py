"""Self-test the grader against the ground truth.

1. IDENTITY: grade {master, Deal-B GT, tracker P8 GT} -> must be ~100%.
2. TRACKER PER-PHASE: grade each snapshot vs its phase -> must be 100% on tracker rows.
3. LAYOUT ROBUSTNESS: synthesize an 'agent-style' workbook with anchors at DIFFERENT
   positions + rounded values + shuffled rent-roll rows -> must still be ~100% within
   tolerance (proves the grader finds values by label, not fixed cell).
"""
import json, os, shutil, tempfile
import openpyxl
import extract as E
import grade as G

HERE = os.path.dirname(os.path.abspath(__file__))
GT = os.path.join(os.path.dirname(os.path.dirname(HERE)), "ground_truth")
MASTER = os.path.join(GT, "Master Model Output", "SouthPark_Centre_Model (Master).xlsx")
DEALB = os.path.join(GT, "Deal B Screening", "Ground_Truth_Deal_B_WACC_Screen.xlsx")
TRK = os.path.join(GT, "Pipeline Tracker")
KEY = G.KEY


def _mkdir():
    d = tempfile.mkdtemp(prefix="sp_sub_")
    return d


def identity_test():
    d = _mkdir()
    shutil.copyfile(MASTER, os.path.join(d, KEY["workbook"]))
    shutil.copyfile(DEALB, os.path.join(d, KEY["dealb_file"]))
    shutil.copyfile(os.path.join(TRK, KEY["tracker"]["8"]["snapshot"]), os.path.join(d, KEY["tracker_file"]))
    print("########## 1. IDENTITY SELF-TEST (GT as submission) ##########")
    ok, n = G.report(G.grade(d, "8"))
    shutil.rmtree(d)
    return ok, n


def tracker_phase_test():
    print("\n########## 2. TRACKER PER-PHASE SELF-TEST ##########")
    allok = True
    for ph in ["1", "2", "3", "6", "8"]:
        d = _mkdir()
        # only the tracker matters here; copy snapshot as the tracker file
        shutil.copyfile(os.path.join(TRK, KEY["tracker"][ph]["snapshot"]), os.path.join(d, KEY["tracker_file"]))
        res = [r for r in G.grade(d, ph) if r[1].startswith("tracker")]
        ok = sum(1 for r in res if r[2]); n = len(res)
        flag = "OK" if ok == n else "FAIL"
        print(f"  P{ph} tracker: {ok}/{n} [{flag}]  " + " | ".join(f"{cid.split('[')[1][:-1]}={got}" for _, cid, _, got, _ in res))
        allok = allok and ok == n
        shutil.rmtree(d)
    return allok


def _synth_submission(errors=None):
    """Build an 'agent-style' submission at a different layout. `errors` overrides:
    anchors[id]=val, rent_rate[name]=val, comp_type[id]=val, calc[(label,month)]=val."""
    errors = errors or {}
    import random
    d = _mkdir()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    # group anchors by tab; write label in col A, value in col C (master used D/B) + junk col B
    by_tab = {}
    for a in KEY["anchors"]:
        by_tab.setdefault(a["tab"], []).append(a)
    for tab, anchors in by_tab.items():
        ws = wb.create_sheet(tab)
        ws["A1"] = f"{tab} (agent layout)"
        r = 3
        for a in anchors:
            exp = a["expected"]
            # round numeric to fewer decimals to exercise tolerance
            if a["kind"] == "num":
                try:
                    fexp = float(exp)
                    # round within tolerance: rates to 4dp, dollars to nearest dollar
                    exp = round(fexp, 4) if abs(fexp) < 10 else float(round(fexp))
                except (TypeError, ValueError):
                    pass
            if a["id"] in errors.get("anchors", {}):
                exp = errors["anchors"][a["id"]]
            ws.cell(r, 1, a["label"])      # label col A
            ws.cell(r, 2, "")              # junk gap col B
            ws.cell(r, 3, exp)             # value col C (different from master)
            r += 1
    # ALL per-item tables, generically, at shifted positions (rent roll incl. benchmark fields,
    # sales+lease comps, leasing buckets). Error overrides: rent_rate / comp_type.
    random.seed(7)
    for pid, ri in KEY["per_item"].items():
        tab = ri["tab"]
        ws = wb[tab] if tab in wb.sheetnames else wb.create_sheet(tab)
        start = (ws.max_row + 3) if (ws.max_row and ws.max_row > 1) else 2
        cols = list(ri["want"].items())                       # (outkey, header text)
        for ci, (_outk, htext) in enumerate(cols, start=1):
            ws.cell(start, ci, htext)
        rows = list(ri["rows"]); random.shuffle(rows)
        for rj, row in enumerate(rows, start=start + 1):
            for ci, (outk, _h) in enumerate(cols, start=1):
                val = row.get(outk)
                if pid == "rent_roll" and outk == "rent":
                    val = errors.get("rent_rate", {}).get(row.get("name"), val)
                if pid in ("sales_comps", "lease_comps") and outk == "ctype":
                    val = errors.get("comp_type", {}).get(row.get("id"), val)
                ws.cell(rj, ci, val)
    # Leverage Options matrix: 'Term' header row with lenders across, metric rows below
    lm = KEY["matrices"]["leverage"]
    ws = wb[lm["tab"]] if lm["tab"] in wb.sheetnames else wb.create_sheet(lm["tab"])
    lr = (ws.max_row + 3) if (ws.max_row and ws.max_row > 1) else 2
    lenders = lm["col_anchors"]
    ws.cell(lr, 1, "Term")
    lcol = {ld: 2 + i for i, ld in enumerate(lenders)}
    for ld, c in lcol.items():
        ws.cell(lr, c, ld)
    metrics = {}
    for kstr, val in lm["cells"].items():
        metric, lender = kstr.split("||"); metrics.setdefault(metric, {})[lender] = val
    for mj, (metric, vals) in enumerate(metrics.items(), start=lr + 1):
        ws.cell(mj, 1, metric)
        for ld, v in vals.items():
            ws.cell(mj, lcol[ld], v)
    # Sensitivity: drop the unique result strings somewhere on Assumptions (presence check)
    sm = KEY["matrices"]["sensitivity"]
    ws = wb[sm["tab"]] if sm["tab"] in wb.sheetnames else wb.create_sheet(sm["tab"])
    sr = (ws.max_row + 3) if (ws.max_row and ws.max_row > 1) else 2
    for i, v in enumerate(sm["cells"].values()):
        ws.cell(sr + i, 1, v)
    # Reconstruct EVERY calc-grid tab at a DIFFERENT layout (shifted header row + columns,
    # block in col A, label in col B) -> proves full-grid (block,label,occ,col) addressing is
    # layout-independent. errors['cell'][(tab, keystr)] overrides a specific cell.
    import json as _json
    cell_err = errors.get("cell", {})
    for tab, spec in KEY.get("calc_grids", {}).items():
        # reuse the sheet if anchors already created it (e.g. Waterfall), append grid below
        ws = wb[tab] if tab in wb.sheetnames else wb.create_sheet(tab)
        row_base = (ws.max_row + 3) if ws.max_row and ws.max_row > 1 else 1
        ws.cell(row_base, 1, f"{tab} (agent layout)")
        # group cells by (block,label,occ) preserving first-seen order
        rows = {}
        for kstr, val in spec["cells"].items():
            block, label, occ, ck = _json.loads(kstr)
            rows.setdefault((block, label, occ), []).append((tuple(ck), val, kstr))
        months = sorted({ck[1] for cells in rows.values() for ck, _, _ in cells if ck[0] == "M"})
        years = sorted({ck[1] for cells in rows.values() for ck, _, _ in cells if ck[0] == "Y"})
        heads = sorted({ck[0] for cells in rows.values() for ck, _, _ in cells if len(ck) == 1})
        base = 4                                   # shift: month strip starts at col 4
        mcol = {m: base + i for i, m in enumerate(months)}
        ycol = {y: base + len(months) + 2 + i for i, y in enumerate(years)}
        hcol = {h: base + len(months) + len(years) + 4 + i for i, h in enumerate(heads)}
        hdr_row = row_base + 1                      # shifted header row (below any anchor content)
        for m, c in mcol.items():
            ws.cell(hdr_row, c, m)                 # consecutive monthly run
        for y, c in ycol.items():
            ws.cell(hdr_row, c, y * 12)            # step-12 annual run
        for h, c in hcol.items():
            ws.cell(hdr_row, c, h)
        r = hdr_row + 1
        for (block, label, occ), cells in rows.items():
            ws.cell(r, 1, block if block not in (None, "None") else "block")
            ws.cell(r, 2, label)
            for ck, val, kstr in cells:
                v = cell_err.get((tab, kstr), val)
                if ck[0] == "M":
                    ws.cell(r, mcol[ck[1]], v)
                elif ck[0] == "Y":
                    ws.cell(r, ycol[ck[1]], v)
                elif len(ck) == 1:
                    ws.cell(r, hcol[ck[0]], v)
                # ('c', n) scalar fallback cells omitted (re-key ambiguously across layouts)
            r += 1
    wb.save(os.path.join(d, KEY["workbook"]))
    # reuse GT for Deal B + tracker
    shutil.copyfile(DEALB, os.path.join(d, KEY["dealb_file"]))
    shutil.copyfile(os.path.join(TRK, KEY["tracker"]["8"]["snapshot"]), os.path.join(d, KEY["tracker_file"]))
    return d


def robustness_test():
    print("\n########## 3. LAYOUT-ROBUSTNESS SELF-TEST (synthetic agent-style workbook) ##########")
    d = _synth_submission()
    ok, n = G.report(G.grade(d, "8"))
    shutil.rmtree(d)
    return ok, n


def detection_test():
    """Inject specific errors; confirm the grader fails EXACTLY those checks (proves teeth)."""
    print("\n########## 4. ERROR-DETECTION SELF-TEST (inject known-wrong values) ##########")
    import json as _json
    def find_cell(tab, label, m, nonzero=False):
        for k, v in KEY["calc_grids"][tab]["cells"].items():
            b, l, o, ck = _json.loads(k)
            if l == label and ck == ["M", m] and (not nonzero or abs(float(v)) > 1000):
                return k
        return None
    pcf_noi = find_cell("Property Cash Flow", "Net Operating Income", 20)          # hardcoded wrong monthly NOI
    pcf_lcf = find_cell("Property Cash Flow", "Total Levered Cash Flow", 30)        # breaks LCF=UCF+FIN @ m30
    # a NON-ZERO per-tenant engine cell (month 1, before Meridian vacates) -> proves per-cell teeth
    rrc_rent = (find_cell("Rent Roll Calc", "Net Base Rent", 1, nonzero=True)
                or find_cell("Rent Roll Calc", "In-Place Rent", 1, nonzero=True))
    injected = {
        "anchors": {"selected_lender": "Atlantic Life Insurance",   # the stale-lender trap
                    "recommendation": "NO-GO", "lp_irr": 0.05},
        "rent_rate": {"Brightpath Insurance": 21.25},                # the un-reconciled trap value
        "comp_type": {3: 9},                                          # misclassified comp (GT=1)
        "cell": {("Property Cash Flow", pcf_noi): 999999.0,
                 ("Property Cash Flow", pcf_lcf): 12345.0,
                 ("Rent Roll Calc", rrc_rent): 0.0},
    }
    expected_fail_ids = {
        "selected_lender", "recommendation", "lp_irr",
        "rent_roll:Brightpath Insurance:rent", "sales_comps:3:ctype",
        f"cell:Property Cash Flow:{pcf_noi}",
        f"cell:Property Cash Flow:{pcf_lcf}",
        f"cell:Rent Roll Calc:{rrc_rent}",
        "invariant:NOI = EGI + Vacancy + OpEx + MgmtFee[m20]",   # NOI@20 now inconsistent
        "invariant:LeveredCF = UnleveredCF + FinancingCF[m30]",  # LCF@30 now inconsistent
    }
    d = _synth_submission(injected)
    res = G.grade(d, "8")
    shutil.rmtree(d)
    failed = {cid for _, cid, ok, _, _ in res if not ok}
    caught = expected_fail_ids & failed
    missed = expected_fail_ids - failed
    unexpected = failed - expected_fail_ids
    print(f"  injected {len(expected_fail_ids)} errors; grader caught {len(caught)}")
    for cid in sorted(expected_fail_ids):
        print(f"    {'CAUGHT' if cid in failed else 'MISSED'}  {cid}")
    if unexpected:
        print(f"  (also failed {len(unexpected)} other checks — collateral from injected breaks, e.g. month-20/30 value mismatches)")
    return len(missed) == 0


if __name__ == "__main__":
    o1, n1 = identity_test()
    t_ok = tracker_phase_test()
    o3, n3 = robustness_test()
    d_ok = detection_test()
    print("\n================ SELF-TEST SUMMARY ================")
    print(f"  Identity:      {o1}/{n1} ({100*o1/n1:.1f}%)")
    print(f"  Tracker/phase: {'PASS' if t_ok else 'FAIL'}")
    print(f"  Robustness:    {o3}/{n3} ({100*o3/n3:.1f}%)")
    print(f"  Detection:     {'PASS (all injected errors caught)' if d_ok else 'FAIL (missed an injected error)'}")
